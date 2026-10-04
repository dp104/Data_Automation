"""Projects of a type (products, jobs, properties, hotels, flights): sites, runs, price history, export."""
import datetime as dt
import json
import os
import re
import sqlite3
import threading
import time
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import project_types as pt
import template_engine as te
import type_engine as tx
from crawler import Fetcher, TypeCrawler, load_robots, normalise

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("UNISCRAPE_DATA_DIR") or HERE
PROJECTS = os.path.join(DATA_DIR, "projects")
RUNS = os.path.join(DATA_DIR, "cache", "projects")
DB = os.path.join(DATA_DIR, "cache", "history.db")
_db_lock = threading.Lock()


# ---------------------------------------------------------------- storage
def pid_of(name):
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:60] or "project"


def save(p):
    os.makedirs(PROJECTS, exist_ok=True)
    with open(os.path.join(PROJECTS, f"{p['id']}.json"), "w") as f:
        json.dump(p, f, indent=1)


def load(pid):
    with open(os.path.join(PROJECTS, re.sub(r"[^a-z0-9-]", "", pid) + ".json")) as f:
        return json.load(f)


def run_file(pid):
    return os.path.join(RUNS, f"{pid}.json")


def last_run(pid):
    try:
        return json.load(open(run_file(pid)))
    except (OSError, ValueError):
        return None


def list_projects():
    out = []
    if os.path.isdir(PROJECTS):
        for n in sorted(os.listdir(PROJECTS)):
            if not n.endswith(".json"):
                continue
            try:
                p = json.load(open(os.path.join(PROJECTS, n)))
            except (OSError, ValueError):
                continue
            r = last_run(p["id"]) or {}
            out.append({"id": p["id"], "name": p["name"], "type": p["type"], "sites": p["sites"],
                        "last_run": r.get("finished_at", ""), "rows": len(r.get("rows", [])),
                        "changes": sum(1 for x in r.get("rows", []) if x.get("_change") in ("up", "down"))})
    return sorted(out, key=lambda x: x["last_run"], reverse=True)


# ---------------------------------------------------------------- price history
def db():
    os.makedirs(os.path.dirname(DB), exist_ok=True)
    con = sqlite3.connect(DB)
    con.execute("create table if not exists snap (project text, item text, name text, price real, currency text, at text)")
    con.execute("create index if not exists snap_i on snap(project, item, at)")
    return con


PRICE_KEY = {"product": "price", "hotels": "night", "flights": "fare", "realestate": "price"}
ITEM_LABEL = {"product": "products", "realestate": "properties", "jobs": "job listings", "flights": "flights",
              "hotels": "hotels", "university": "course pages"}


def apply_history(p, rows, at):
    """Compare each row's price with the last check of the same item; store this check."""
    key = PRICE_KEY.get(p["type"])
    if not key:
        return rows
    with _db_lock:
        con = db()
        for r in rows:
            item = (r.get("url") or r.get("name") or r.get("title") or "")[:500]
            price = r.get(key)
            prev = con.execute("select price, at from snap where project=? and item=? order by at desc limit 1",
                               (p["id"], item)).fetchone()
            if prev is None:
                r["_change"], r["change"] = "new", "First check"
            elif price is None or prev[0] is None:
                r["_change"], r["change"] = "unknown", "No price on page" if price is None else "First price"
            elif abs(float(price) - float(prev[0])) < 0.005:
                r["_change"], r["change"] = "same", "No change"
            else:
                diff = float(price) - float(prev[0])
                pct = diff / float(prev[0]) * 100 if prev[0] else 0
                r["_change"] = "up" if diff > 0 else "down"
                r["change"] = f"{'Up' if diff > 0 else 'Down'} {abs(pct):.1f}% (was {prev[0]:g} on {prev[1][:10]})"
                r["_was"] = prev[0]
            if price is not None:
                con.execute("insert into snap values (?,?,?,?,?,?)", (p["id"], item, r.get("name") or r.get("title"), price,
                                                                      r.get("currency"), at))
        con.commit()
        con.close()
    return rows


def history(pid, item):
    with _db_lock:
        con = db()
        rows = con.execute("select price, currency, at from snap where project=? and item=? order by at", (pid, item)).fetchall()
        con.close()
    return [{"price": a, "currency": b, "at": c} for a, b, c in rows]


# ---------------------------------------------------------------- collecting
JOBS = {}


def run(pid, max_pages=1500, log=print):
    """Collect every site of a project. Rows keep the best data per item (item page beats listing card)."""
    p = load(pid)
    t = pt.TYPES[p["type"]]
    job = JOBS[pid]
    items, order = {}, []
    notes = []
    for site in p["sites"]:
        fetcher = Fetcher()
        try:
            robots = load_robots(fetcher, site)
            if robots is not None and not robots.can_fetch("*", site):
                notes.append(f"{site}: the site's robots rules do not allow this page to be collected, so it was skipped.")
                log(notes[-1])
                continue
            st0, _, txt0, challenge = fetcher.http_get(site)
            if challenge:  # a browser check page (e.g. "Just a moment..."), not a refusal: open it like a browser
                log(f"{site}: site shows a browser check -> using headless Chrome")
                fetcher.browser_only = True
            elif st0 in (401, 403, 429):
                notes.append(f"{site}: the site refused access (HTTP {st0}). It may be limiting requests; try again later.")
                log(notes[-1])
                continue
            card_urls = set()
            listing_paths = {urlparse(site).path.rstrip("/")}

            def on_page(u, html, _site=site):
                key = normalise(u.split("#")[0])
                path = urlparse(u).path.rstrip("/")
                rows, kind = tx.extract_page(t, html, u, expect_item=key in card_urls)
                if kind == "item" and path in listing_paths and key not in card_urls:
                    # same address as a listing, only "?location=..." differs: a filtered listing, not an item
                    from bs4 import BeautifulSoup
                    rows, kind = tx.card_rows(t, BeautifulSoup(html, "lxml"), u, tx.site_name(None, u)), "list"
                if kind == "list":
                    listing_paths.add(path)
                if kind == "item" and rows:
                    r = rows[0]
                    k = normalise((r.get("url") or u).split("#")[0]) or key
                    old = items.get(k, {})
                    merged = {**old, **{a: b for a, b in r.items() if b not in (None, "")}}
                    merged["_site"] = _site
                    if k not in items:
                        order.append(k)
                    items[k] = merged
                    crawler.note_item(u)
                    job["count"] = len(items)
                    return True
                if kind == "list":
                    from bs4 import BeautifulSoup
                    if key in card_urls:
                        return False  # an item's own page: its "related"/"benefits" blocks are not more items
                    before = len(items)
                    for r in rows:
                        link = r.get("url")
                        k = normalise(link.split("#")[0]) if link else None
                        if not k:
                            continue
                        h = urlparse(k).netloc.lower()
                        if not (h == crawler.dom or h.endswith("." + crawler.dom)):
                            continue  # a card linking to another website is not an item of this site
                        if urlparse(k).path.rstrip("/") in listing_paths:
                            crawler.push_boost(k, 2, 1)   # a filter/sort link of the listing: maybe more results,
                            continue                      # but never an item itself
                        if k not in items:
                            items[k] = dict(r, _site=_site, _from="listing")
                            order.append(k)
                        if True:
                            card_urls.add(k)
                            crawler.push_boost(k, 0.1, 1)       # open each item's own page
                            crawler.note_item(k)
                    if len(items) > before:  # only page further while pages keep bringing new items
                        for nxt in tx.next_page_links(BeautifulSoup(html, "lxml"), u):
                            crawler.push_boost(nxt, 0.05, 1)    # follow "next page"
                    job["count"] = len(items)
                    return bool(rows)
                return False

            crawler = TypeCrawler(site, fetcher, on_page, max_pages=max_pages, log=log, robots=robots,
                                  item_label=ITEM_LABEL.get(p["type"], "items"))
            crawler.run([site])
            if getattr(fetcher, "refusals", 0):
                notes.append(f"{site}: the site asked us to slow down {fetcher.refusals} times; the tool waited and slowed down. "
                             "If rows are missing, collect again later.")
            if crawler.blocked:
                notes.append(f"{site}: {crawler.blocked} links were skipped because the site's robots rules do not allow them.")
        finally:
            fetcher.close()
    rows = [items[k] for k in order]
    at = dt.datetime.now().isoformat(timespec="seconds")
    for r in rows:
        r["checked"] = at[:16].replace("T", " ")
    rows = apply_history(p, rows, at)
    os.makedirs(RUNS, exist_ok=True)
    with open(run_file(pid), "w") as f:
        json.dump({"finished_at": at, "rows": rows, "notes": notes}, f)
    return rows, notes


# ---------------------------------------------------------------- formats
def standard_template(type_id):
    """The type's standard format as a sheet template (created once)."""
    t = pt.TYPES[type_id]
    tid = f"standard-{type_id}"
    path_json = os.path.join(te.TEMPLATES, f"{tid}.json")
    if os.path.exists(path_json):
        return te.load(tid)
    os.makedirs(te.TEMPLATES, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = t["name"][:31]
    thin = Side(style="thin", color="FFFFFF")
    for i, f in enumerate(t["fields"], 1):
        c = ws.cell(1, i, f["header"])
        c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
        c.fill = PatternFill("solid", fgColor="0D6B5E")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(left=thin, right=thin)
        ws.column_dimensions[c.column_letter].width = {"url": 60, "text": 28, "money": 14, "date": 16}.get(f["type"], 12) \
            if f["key"] not in ("name", "title") else 48
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    xlsx = os.path.join(te.TEMPLATES, f"{tid}.xlsx")
    wb.save(xlsx)
    from openpyxl.utils import get_column_letter
    tpl = {"id": tid, "name": f"Standard: {t['name']}", "kind": "standard", "type": type_id, "answered": True,
           "sample": os.path.basename(xlsx), "header_row": 1, "sheet_title": ws.title, "per_intake": False,
           "sheet_style": None, "rows": [], "created": dt.date.today().isoformat(),
           "columns": [{"col": i, "letter": get_column_letter(i), "header": f["header"], "example": "", "fixed": "",
                        "field": "serial" if f.get("special") == "serial" else "typefield", "key": f["key"]}
                       for i, f in enumerate(t["fields"], 1)]}
    te.save(tpl)
    return tpl


def suggest_mapping(tpl, type_id):
    """For the user's own sheet: which type field each column most likely is (by header)."""
    from difflib import SequenceMatcher
    t = pt.TYPES[type_id]
    out = []
    for c in tpl["columns"]:
        h = c["header"].lower()
        best, score = None, 0
        for f in t["fields"]:
            words = [f["header"].lower(), f["key"]] + [l.lower() for l in f.get("labels", [])]
            sc = max(SequenceMatcher(None, h, w).ratio() for w in words)
            if any(w in h for w in words if len(w) > 3):
                sc = max(sc, 0.8)
            if sc > score:
                best, score = f["key"], sc
        out.append({"header": c["header"], "key": best if score >= 0.6 else None,
                    "serial": c.get("field") == "serial", "formula": bool(c.get("formula"))})
    return out


def export(p, out_dir, fixed=None):
    r = last_run(p["id"])
    if not r or not r["rows"]:
        raise ValueError("Collect the data first.")
    if p.get("format", "standard") == "standard":
        tpl = standard_template(p["type"])
        mapping = {c["header"]: c.get("key") for c in tpl["columns"]}
    else:
        tpl = te.load(p["format"])
        mapping = {m["header"]: m.get("key") for m in p.get("mapping") or []}
    rows = [{h: row.get(k) if k else None for h, k in mapping.items()} for row in r["rows"]]
    t = pt.TYPES[p["type"]]
    safe = re.sub(r'[\\/:*?"<>|]+', " ", p["name"]).strip()[:80] or "Export"
    stamp = time.strftime("%Y-%m-%d")
    out = os.path.join(out_dir, f"{safe} - {stamp}.xlsx")
    n = 2
    while os.path.exists(out):
        out = os.path.join(out_dir, f"{safe} - {stamp} ({n}).xlsx")
        n += 1
    te.render_rows(tpl, rows, out, fixed or {})
    key_cols = [f["key"] for f in t["fields"] if f.get("special") not in ("serial", "site", "page_url", "checked", "change", "computed")]
    missing = sum(1 for x in r["rows"] if not any(x.get(k) for k in tx.KEY_FIELDS.get(p["type"], key_cols[:1])))
    return {"xlsx": os.path.basename(out), "rows": len(rows), "missing_key": missing}
