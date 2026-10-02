"""Writes courses in the Flyurdream university sheet format (one sheet per intake month)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color

HEADERS = ["Sl.No", "Country ", "University Link", "University", "Campus", "City", "Region", "Level",
           "Mode Of Education", "Course", "Course Link", "Duration", "First Year fee", "Scholarship", "Intake",
           "Cost of living", "Minimum Deposit", "Credibilty interview", "Application Fee"]
WIDTHS = [5.44, 13.89, 20.55, 22.0, 15.55, 8.78, 8.0, 13.33, 17.55, 72.55, 108.0, 8.33, 11.66, 10.66, 9.78,
          11.78, 16.11, 17.44, 14.0]

MONTH_SHORT = {"January": "Jan", "February": "Feb", "March": "Mar", "April": "Apr", "May": "May", "June": "Jun",
               "July": "Jul", "August": "Aug", "September": "Sep", "October": "Oct", "November": "Nov",
               "December": "Dec"}


def write_workbook(path, sheets, uni):
    """sheets: list of (month_full_name, [records]); uni: dict of fixed per-university values."""
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin")
    head_font = Font(name="Calibri", size=11, bold=True, color=Color(theme=0))
    head_fill = PatternFill("solid", fgColor=Color(theme=1))
    head_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center")
    for month, recs in sheets:
        ws = wb.create_sheet(MONTH_SHORT.get(month, month))
        for c, h in enumerate(HEADERS, 1):
            cell = ws.cell(1, c, h)
            cell.font, cell.fill, cell.border = head_font, head_fill, head_border
            cell.alignment = Alignment(horizontal="left" if h == "Level" else "center")
        for i, w in enumerate(WIDTHS):
            ws.column_dimensions[chr(65 + i)].width = w
        for i, r in enumerate(recs, 1):
            row = i + 1
            level = r["level_out"]
            values = [i, uni["country"], uni["url"], uni["name"], r.get("campus_out") or uni["campus"], uni["city"],
                      uni["region"] or None, level, r["mode"], r["course"], r["url"], r["duration"], r["fee"],
                      uni["scholarship"].get(level, uni["scholarship"].get("*", 0)), month, uni["cost_of_living"],
                      f"=M{row}*{uni['deposit_pct']}", uni["credibility"], uni["application_fee"]]
            for c, v in enumerate(values, 1):
                ws.cell(row, c, v)
            ws.cell(row, 3).hyperlink = uni["url"]
            ws.cell(row, 3).style = "Hyperlink"
            ws.cell(row, 11).hyperlink = r["url"]
            ws.cell(row, 11).style = "Hyperlink"
            ws.cell(row, 17).number_format = "0"
    wb.save(path)
