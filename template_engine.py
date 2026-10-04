"""Output formats ("templates") learned from a sample Excel the team uploads.

learn(sample.xlsx)  -> template: the sample's exact columns (order, header text), what each column is filled
                       from, how values are written (e.g. "September" vs "Sep 2026", 12 vs "1 Year"),
                       formulas (=M2*0.34), sheet layout (one sheet per intake or one sheet) and styling.
render(template, ...) writes scanned courses in exactly that format. A column the scan cannot fill is
left blank (or gets the fixed value the user typed) - values are never invented.
"""
import copy
import datetime as dt
import json
import os
import re
import shutil

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("UNISCRAPE_DATA_DIR") or HERE
TEMPLATES = os.path.join(DATA_DIR, "templates")


def seed_standard_format():
    """Copy the built-in standard-university format into TEMPLATES if it isn't there yet - needed
    the first time TEMPLATES lives outside the repo (UNISCRAPE_DATA_DIR set, e.g. a Render disk)."""
    if os.path.abspath(TEMPLATES) == os.path.abspath(os.path.join(HERE, "templates")):
        return
    os.makedirs(TEMPLATES, exist_ok=True)
    for name in ("standard-university.json", "standard-university.xlsx"):
        dst = os.path.join(TEMPLATES, name)
        src = os.path.join(HERE, "templates", name)
        if not os.path.exists(dst) and os.path.exists(src):
            shutil.copyfile(src, dst)

MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]
MON3 = {m[:3].lower(): m for m in MONTHS}

# What a column can be filled from. "uni" fields are per-university values typed on the details form.
FIELDS = {
    "serial":          ("Row number (1, 2, 3...)", "row"),
    "university_name": ("University name", "uni"),
    "university_url":  ("University website link", "uni"),
    "country":         ("Country", "uni"),
    "campus":          ("Campus", "row"),
    "city":            ("City", "uni"),
    "region":          ("Region / state", "uni"),
    "level":           ("Education level", "row"),
    "mode":            ("Mode of education", "row"),
    "attendance":      ("Full-time / part-time", "row"),
    "course":          ("Course name", "row"),
    "course_url":      ("Course page link", "row"),
    "duration":        ("Duration", "row"),
    "fee":             ("International first-year fee", "row"),
    "currency":        ("Currency", "uni"),
    "scholarship":     ("Scholarship (per level)", "uni"),
    "intake":          ("Intake / start month", "row"),
    "cost_of_living":  ("Cost of living", "uni"),
    "deposit":         ("Minimum deposit", "uni"),
    "credibility":     ("Credibility interview", "uni"),
    "application_fee": ("Application fee", "uni"),
    "fixed":           ("Same fixed value on every row", "fixed"),
    "blank":           ("Leave blank", "blank"),
}

SYNONYMS = [
    ("serial", r"^(s\.?\s?no\.?|sl\.?\s?no\.?|sr\.?\s?no\.?|serial( no\.?| number)?|#|no\.?|sno|slno)$"),
    ("university_url", r"(university|uni|institution|college)\s*(link|url|website|site)|^website$"),
    ("course_url", r"(course|programme|program)\s*(link|url|page)|^link$|^url$"),
    ("university_name", r"^(university|uni|institution|college|school)(\s*name)?$"),
    ("country", r"^country"),
    ("campus", r"^campus|^location$|^study location"),
    ("city", r"^city"),
    ("region", r"^(region|state|province|county)"),
    ("level", r"level|^degree( type)?$|qualification"),
    ("mode", r"mode|delivery|study type"),
    ("attendance", r"attendance|full[\s-]?time.*part|study pattern|study load"),
    ("course", r"^(course|programme|program)(\s*(name|title))?$|^course\s*name|^programme\s*name"),
    ("duration", r"duration|length|^years?$|^months?$"),
    ("scholarship", r"scholarship|bursary"),
    ("intake", r"intake|start\s*(date|month)?|entry\s*(date|month|point)"),
    ("cost_of_living", r"living"),
    ("deposit", r"deposit"),
    ("credibility", r"credib|interview"),
    ("application_fee", r"application\s*fee|app\.?\s*fee"),
    ("currency", r"^currency|^curr\.?$"),
    ("fee", r"fee|tuition|cost|price"),
]

LEVEL_CLASSES = [
    ("DBA", r"\bdba\b|business administration doctor"),
    ("Research", r"research|ph\.?d|doctor|mphil|mres"),
    ("Foundation", r"foundation|pathway|pre-?master|pre-?sessional"),
    ("Undergraduate", r"under\s*grad|\bug\b|bachelor|bsc|\bba\b|degree"),
    ("Postgraduate", r"post\s*grad|\bpg\b|master|msc|\bma\b|mba|(?<!under)graduate"),
    ("Diploma/Certificate", r"diploma|certificate|hnd|hnc"),
    ("Short course/Other", r"short|cpd|other"),
]


def norm(s):
    return re.sub(r"[\s_]+", " ", str(s or "")).strip().lower()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-") or "format"


# ---------------------------------------------------------------- learning
def find_header(ws):
    for r in range(1, min(ws.max_row, 15) + 1):
        vals = [c.value for c in ws[r]]
        if sum(1 for v in vals if isinstance(v, str) and v.strip()) >= 3:
            return r
    return None


def month_of(v):
    if isinstance(v, (dt.date, dt.datetime)):
        return MONTHS[v.month - 1]
    m = re.match(r"\s*([A-Za-z]{3,9})", str(v or ""))
    return MON3.get(m.group(1)[:3].lower()) if m and m.group(1)[:3].lower() in MON3 else None


def intake_style(samples):
    """How intake values are written in the sample."""
    for v in samples:
        if isinstance(v, (dt.date, dt.datetime)):
            return "date"
        s = str(v).strip()
        if re.fullmatch(r"[A-Za-z]{3}\s*[-'/ ]\s*\d{2}", s):
            return "Mon-YY" if "-" in s else "Mon YY"
        if re.fullmatch(r"[A-Za-z]{4,9}\s+\d{4}", s):
            return "Month YYYY"
        if re.fullmatch(r"[A-Za-z]{3}\s+\d{4}", s):
            return "Mon YYYY"
        if re.fullmatch(r"[A-Za-z]{3}\.?", s):
            return "Mon"
        if re.fullmatch(r"[A-Za-z]{4,9}", s):
            return "Month"
    return "Month"


def duration_style(samples):
    for v in samples:
        if isinstance(v, (int, float)):
            return "years_number" if all(isinstance(x, (int, float)) and x <= 6 for x in samples) and \
                any(isinstance(x, float) and not float(x).is_integer() for x in samples) else "months_number"
        s = str(v).lower()
        if "month" in s:
            return "months_text"
        if "year" in s or "yr" in s:
            return "years_text"
    return "months_number"


def guess_field(header, samples):
    h = norm(header)
    for field, pat in SYNONYMS:
        if re.search(pat, h):
            if field == "fee" and re.search(r"application", h):
                continue
            return field
    vals = [v for v in samples if v not in (None, "")]
    if vals and all(isinstance(v, str) and v.startswith("http") for v in vals):
        paths = [re.sub(r"^https?://[^/]+", "", v).strip("/") for v in vals]
        return "university_url" if all(not p for p in paths) else "course_url"
    if vals and all(month_of(v) for v in vals):
        return "intake"
    return None


def learn(path, name):
    wb = load_workbook(path)  # formulas kept as text
    data_sheets = []
    for ws in wb.worksheets:
        hr = find_header(ws)
        if hr and ws.max_row > hr:
            data_sheets.append((ws, hr))
    if not data_sheets:
        raise ValueError("No sheet with a header row and data rows was found in this file.")
    ws0, hr = data_sheets[0]
    headers = []
    for c in ws0[hr]:
        if c.value not in (None, "") and str(c.value).strip():
            headers.append((c.column, str(c.value)))
    # sample values per column across all data sheets
    samples = {col: [] for col, _ in headers}
    first_row = {}
    for ws, h in data_sheets:
        for r in range(h + 1, min(ws.max_row, h + 300) + 1):
            if all(ws.cell(r, col).value in (None, "") for col, _ in headers):
                continue
            for col, _ in headers:
                v = ws.cell(r, col).value
                if v not in (None, ""):
                    samples[col].append(v)
                    first_row.setdefault(col, (r, v, ws.cell(r, col)))
    # one sheet per intake?
    per_intake = len(data_sheets) > 1 and all(month_of(ws.title) for ws, _ in data_sheets)
    sheet_style = intake_style([ws.title for ws, _ in data_sheets]) if per_intake else None

    cols = []
    for col, header in headers:
        vals = samples[col]
        field = guess_field(header, vals)
        spec = {"col": col, "letter": get_column_letter(col), "header": header, "field": field or "blank",
                "example": "" if not vals else str(vals[0])[:80], "fixed": ""}
        formulas = [v for v in vals if isinstance(v, str) and v.startswith("=")]
        if formulas:  # keep the sample's formula, re-pointed at each row (=M2*0.34 -> =M{row}*0.34)
            r0 = first_row[col][0]
            spec["formula"] = re.sub(rf"(?<=[A-Z]){r0}\b", "{row}", formulas[0])
            spec["field"] = field or "fixed"
        if spec["field"] == "intake":
            spec["style"] = intake_style(vals)
        elif spec["field"] == "duration":
            spec["style"] = duration_style(vals)
        elif spec["field"] == "level":
            spec["labels"] = learn_labels(vals, LEVEL_CLASSES)
        elif spec["field"] == "mode":
            spec["labels"] = learn_mode_labels(vals)
        elif spec["field"] == "attendance":
            spec["labels"] = learn_attendance_labels(vals)
        elif not field and not formulas:
            distinct = {str(v) for v in vals}
            if len(distinct) == 1:  # same value on every sample row: offer it as a fixed column, user confirms
                spec["field"] = "fixed"
                spec["hint"] = next(iter(distinct))
        cols.append(spec)

    # a few complete sample rows (text), used by "Any website" mode to find these values on example pages
    sample_rows = []
    for r in range(hr + 1, min(ws0.max_row, hr + 25) + 1):
        vals = [ws0.cell(r, col).value for col, _ in headers]
        if any(v not in (None, "") for v in vals):
            sample_rows.append(["" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
                                for v in vals])
    tid = slug(name)
    os.makedirs(TEMPLATES, exist_ok=True)
    sample_copy = os.path.join(TEMPLATES, f"{tid}.xlsx")
    if os.path.abspath(path) != os.path.abspath(sample_copy):
        shutil.copyfile(path, sample_copy)
    return {"id": tid, "name": name, "sample": os.path.basename(sample_copy), "header_row": hr,
            "sheet_title": ws0.title, "per_intake": per_intake, "sheet_style": sheet_style,
            "columns": cols, "rows": sample_rows[:20], "created": dt.date.today().isoformat()}


def learn_labels(vals, classes):
    """Map our level names to the words the sample uses (e.g. Postgraduate -> 'PG')."""
    out = {}
    from collections import Counter
    for v, _ in Counter(str(x).strip() for x in vals).most_common():  # most used wording wins
        for cls, pat in classes:
            if re.search(pat, v.lower()):
                out.setdefault(cls, v)
                break
    return out


def learn_mode_labels(vals):
    out = {}
    for v in {str(x).strip() for x in vals}:
        lv = v.lower()
        if "online" in lv or "distance" in lv:
            out.setdefault("Online", v)
        elif "blend" in lv or "hybrid" in lv:
            out.setdefault("Blended", v)
        else:
            out.setdefault("On campus", v)
    return out


def learn_attendance_labels(vals):
    out = {}
    for v in {str(x).strip() for x in vals}:
        out.setdefault("Part-time" if "part" in v.lower() else "Full-time", v)
    return out


# ---------------------------------------------------------------- questions about the sample
OUR_LEVELS = ["Undergraduate", "Postgraduate", "Research", "DBA", "Foundation", "Diploma/Certificate", "Short course/Other"]
OUR_MODES = ["On campus", "Online", "Blended"]


def questions(t):
    """Things in the sample that are unclear and need the user's answer before scraping."""
    qs = []
    cols = t["columns"]
    for i, c in enumerate(cols):
        ex = c.get("example") or ""
        if c["field"] in ("fixed", "blank") and not c.get("formula"):
            qs.append({"id": f"col{i}", "kind": "column", "col": i, "header": c["header"],
                       "text": f'Your sample has a column "{c["header"]}"' + (f' (example: "{ex}")' if ex else " with no example values") +
                               ". Websites don't list this, so what should go in it?",
                       "default": "fixed" if c["field"] == "fixed" else "blank"})
        elif not ex and c["field"] not in ("serial",):
            qs.append({"id": f"col{i}", "kind": "confirm_field", "col": i, "header": c["header"],
                       "text": f'The column "{c["header"]}" has no example values. Should it contain: {FIELDS[c["field"]][0]}?',
                       "default": c["field"]})
        if c["field"] == "duration" and c.get("style") == "months_number":
            nums = [x for x in [ex] if str(x).replace(".", "").isdigit()]
            if nums and float(nums[0]) <= 6:
                qs.append({"id": f"dur{i}", "kind": "duration_unit", "col": i, "header": c["header"],
                           "text": f'In "{c["header"]}" the example is {ex}. Is that in years or in months?', "default": "years_number"})
        if c["field"] == "fee":
            qs.append({"id": f"fee{i}", "kind": "info_fee", "col": i, "header": c["header"],
                       "text": f'"{c["header"]}" will contain the international (overseas) first-year tuition fee, as a plain number. '
                               "Courses with no published international fee are left empty. Is that right?", "default": "yes"})
        if c.get("formula"):
            qs.append({"id": f"frm{i}", "kind": "formula", "col": i, "header": c["header"],
                       "text": f'"{c["header"]}" uses a formula in your sample ({c["formula"].replace("{row}", "n")}). '
                               "Keep that formula on every row?", "default": "keep"})
        if c["field"] == "level":
            missing = [l for l in OUR_LEVELS if l not in c.get("labels", {})]
            qs.append({"id": f"lvl{i}", "kind": "labels", "col": i, "header": c["header"], "what": "level",
                       "text": f'How should education levels be written in "{c["header"]}"? Your sample uses: '
                               + ", ".join(sorted(set(c.get("labels", {}).values()))) + ".",
                       "labels": {l: c.get("labels", {}).get(l, l) for l in OUR_LEVELS}, "missing": missing})
        if c["field"] == "mode":
            qs.append({"id": f"mod{i}", "kind": "labels", "col": i, "header": c["header"], "what": "mode",
                       "text": f'How should study mode be written in "{c["header"]}"?',
                       "labels": {m: c.get("labels", {}).get(m, m) for m in OUR_MODES},
                       "missing": [m for m in OUR_MODES if m not in c.get("labels", {})]})
    if not t["per_intake"] and any(c["field"] == "intake" for c in cols):
        qs.append({"id": "layout", "kind": "layout", "text": "Your sample is one sheet with an intake column. "
                   "Keep one sheet for all intakes, or make a separate sheet for each intake?", "default": "single"})
    if not any(c["field"] == "course" for c in cols):
        qs.append({"id": "nocourse", "kind": "warning", "text": "No column for the course name was found. "
                   "Pick the right column below, or the sheet will not say which course each row is."})
    return qs


# ---------------------------------------------------------------- storage
def save(t):
    os.makedirs(TEMPLATES, exist_ok=True)
    with open(os.path.join(TEMPLATES, f"{t['id']}.json"), "w") as f:
        json.dump(t, f, indent=1)


def load(tid):
    with open(os.path.join(TEMPLATES, f"{slug(tid)}.json")) as f:
        return json.load(f)


def list_templates():
    out = []
    if os.path.isdir(TEMPLATES):
        for n in sorted(os.listdir(TEMPLATES)):
            if n.endswith(".json"):
                try:
                    t = json.load(open(os.path.join(TEMPLATES, n)))
                    out.append({"id": t["id"], "name": t["name"], "columns": len(t["columns"]),
                                "per_intake": t["per_intake"], "created": t.get("created", "")})
                except (OSError, ValueError, KeyError):
                    pass
    return out


def needs_intake_year(t):
    return any(c["field"] == "intake" and "Y" in c.get("style", "") or c.get("style") == "date"
               for c in t["columns"]) or bool(t.get("per_intake") and "Y" in (t.get("sheet_style") or ""))


def uni_fields(t):
    """Per-university values this format needs on the details form (only what its columns use)."""
    used = []
    for c in t["columns"]:
        kind = FIELDS.get(c["field"], ("", "blank"))[1]
        if kind == "uni" and c["field"] not in used and not c.get("formula"):
            used.append(c["field"])
    return used


# ---------------------------------------------------------------- rendering
def intake_year(month, today=None):
    """Next upcoming occurrence of that intake month."""
    today = today or dt.date.today()
    m = MONTHS.index(month) + 1
    return today.year if m >= today.month else today.year + 1


def fmt_intake(month, style, year=None):
    if month == "Unknown" or month not in MONTHS:
        return ""
    y = int(year) if str(year or "").strip().isdigit() else intake_year(month)
    return {"Month": month, "Mon": month[:3], "Month YYYY": f"{month} {y}", "Mon YYYY": f"{month[:3]} {y}",
            "Mon-YY": f"{month[:3]}-{str(y)[2:]}", "Mon YY": f"{month[:3]} {str(y)[2:]}",
            "date": dt.datetime(y, MONTHS.index(month) + 1, 1)}.get(style, month)


def fmt_duration(months, style):
    if not months:
        return None
    if style == "months_number":
        return months
    if style == "years_number":
        y = months / 12
        return int(y) if float(y).is_integer() else round(y, 1)
    if style == "months_text":
        return f"{months} Months"
    y = months / 12
    if float(y).is_integer():
        return f"{int(y)} Year" + ("s" if y != 1 else "")
    return f"{months} Months"


def cell_value(spec, row, uni, n, month):
    uni = row.get("_uni") or uni  # combined files: each row carries its own university's values
    f = spec["field"]
    if spec.get("formula"):
        return spec["formula"].replace("{row}", str(n))
    if f == "serial":
        return n - spec.get("_first", 2) + 1
    if f == "blank":
        return None
    if f == "fixed":
        v = (uni.get("fixed") or {}).get(spec["header"], "")
        return num_or_text(v)
    if f == "university_name":
        return uni.get("name") or None
    if f == "university_url":
        return uni.get("url") or None
    if f == "campus":
        return row.get("campus_out") or uni.get("campus") or None
    if f == "level":
        return spec.get("labels", {}).get(row["level_out"], row["level_out"])
    if f == "mode":
        return spec.get("labels", {}).get(row["mode"], row["mode"])
    if f == "attendance":
        key = "Part-time" if row.get("part_time") else "Full-time"
        return spec.get("labels", {}).get(key, key)
    if f == "course":
        return row["course"]
    if f == "course_url":
        return row["url"]
    if f == "duration":
        return fmt_duration(row.get("duration"), spec.get("style", "months_number"))
    if f == "fee":
        return row.get("fee") or None
    if f == "intake":
        return fmt_intake(month, spec.get("style", "Month"), uni.get("intake_year")) or None
    if f == "scholarship":
        sch = uni.get("scholarship") or {}
        return num_or_text(sch.get(row["level_out"], ""))
    if f in ("country", "city", "region", "currency", "cost_of_living", "deposit", "credibility", "application_fee"):
        return num_or_text(uni.get(f, ""))
    return None


def num_or_text(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if re.fullmatch(r"-?\d+(\.\d+)?", s.replace(",", "")):
        f = float(s.replace(",", ""))
        return int(f) if f.is_integer() else f
    return s


def sheet_name(month, t):
    if not t["per_intake"]:
        return t["sheet_title"]
    if month == "Unknown":
        return "Unknown"
    st = t.get("sheet_style") or "Mon"
    v = fmt_intake(month, st)
    return v.strftime("%b %Y") if isinstance(v, dt.datetime) else str(v)[:31]


def render(t, sheets, uni, out_path):
    """sheets: [(month, [row dicts])] from uniscrape.build_sheets."""
    sample_wb = load_workbook(os.path.join(TEMPLATES, t["sample"]))
    sws = sample_wb[t["sheet_title"]] if t["sheet_title"] in sample_wb.sheetnames else sample_wb.worksheets[0]
    hr = t["header_row"]
    wb = Workbook()
    wb.remove(wb.active)
    groups = sheets if t["per_intake"] else [(None, [(m, r) for m, rows in sheets for r in rows])]
    for month, rows in groups:
        ws = wb.create_sheet(sheet_name(month, t) if month else t["sheet_title"])
        # rows above the header (titles, notes) and the header itself, with the sample's styling
        for r in range(1, hr + 1):
            ws.row_dimensions[r].height = sws.row_dimensions[r].height
            for c in t["columns"]:
                src = sws.cell(r, c["col"])
                dst = ws.cell(r, c["col"], c["header"] if r == hr else src.value)
                copy_style(src, dst)
        for c in t["columns"]:
            w = sws.column_dimensions[c["letter"]].width
            if w:
                ws.column_dimensions[c["letter"]].width = w
        style_row = hr + 1
        for i, item in enumerate(rows):
            m, row = (month, item) if month else item
            n = hr + 1 + i
            for c in t["columns"]:
                c["_first"] = hr + 1
                v = cell_value(c, row, uni, n, m)
                cell = ws.cell(n, c["col"], v)
                copy_style(sws.cell(style_row, c["col"]), cell)
                if isinstance(v, str) and v.startswith("http"):
                    cell.hyperlink = v
        if t["columns"] and hr > 0 and sws.freeze_panes:
            ws.freeze_panes = sws.freeze_panes
    wb.save(out_path)


def copy_style(src, dst):
    if src.has_style:
        dst.font = copy.copy(src.font)
        dst.fill = copy.copy(src.fill)
        dst.border = copy.copy(src.border)
        dst.alignment = copy.copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy.copy(src.protection)


def render_rows(t, rows, out_path, fixed=None):
    """'Any website' mode: rows are {header: value}. One sheet, the sample's columns and styling only."""
    fixed = fixed or {}
    sample_wb = load_workbook(os.path.join(TEMPLATES, t["sample"]))
    sws = sample_wb[t["sheet_title"]] if t["sheet_title"] in sample_wb.sheetnames else sample_wb.worksheets[0]
    hr = t["header_row"]
    wb = Workbook()
    ws = wb.active
    ws.title = t["sheet_title"][:31]
    for r in range(1, hr + 1):
        ws.row_dimensions[r].height = sws.row_dimensions[r].height
        for c in t["columns"]:
            src = sws.cell(r, c["col"])
            copy_style(src, ws.cell(r, c["col"], c["header"] if r == hr else src.value))
    for c in t["columns"]:
        w = sws.column_dimensions[c["letter"]].width
        if w:
            ws.column_dimensions[c["letter"]].width = w
    for i, row in enumerate(rows):
        n = hr + 1 + i
        for c in t["columns"]:
            h = c["header"]
            if c.get("formula"):
                v = c["formula"].replace("{row}", str(n))
            elif row.get(h) not in (None, ""):
                v = row[h]
            elif c["field"] == "serial":
                v = i + 1
            else:
                v = num_or_text(fixed.get(h, "")) if fixed.get(h, "") != "" else None
            cell = ws.cell(n, c["col"], v)
            copy_style(sws.cell(hr + 1, c["col"]), cell)
            if isinstance(v, str) and v.startswith("http"):
                cell.hyperlink = v
    if sws.freeze_panes:
        ws.freeze_panes = sws.freeze_panes
    wb.save(out_path)
